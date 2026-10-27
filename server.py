"""정책평가팀 HR관리앱 로컬 서버.

엑셀(연차 일정, 주간업무계획, 사전협의 접수·검토 현황)을 요청 시마다 읽어 JSON으로 제공하고 static/ 화면을 서빙한다.
파일 수정시각이 바뀌지 않으면 캐시를 재사용한다. 공지·웹 입력은 화면에서 Supabase로 직접 읽고 쓴다.
"""
import glob
import json
import re
import threading
import unicodedata
import webbrowser
from datetime import date, datetime, timedelta
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

BASE = Path(__file__).resolve().parent
STATIC = BASE / "static"
_lock = threading.Lock()
_cache = {}


def load_config():
    with open(BASE / "config.json", encoding="utf-8") as f:
        return json.load(f)


def cached(path, parser):
    """파일 mtime 기준 캐시. 읽기 실패 시 직전 결과와 오류 메시지를 함께 돌려준다."""
    p = Path(path)
    try:
        mtime = p.stat().st_mtime
    except OSError as e:
        return None, f"파일을 찾을 수 없습니다: {p} ({e})"
    hit = _cache.get(str(p))
    if hit and hit[0] == mtime:
        return hit[1], None
    try:
        data = parser(p)
    except Exception as e:  # 엑셀이 열려 있거나 저장 중인 경우 등
        return (hit[1] if hit else None), f"엑셀을 읽지 못했습니다: {p.name} ({e})"
    _cache[str(p)] = (mtime, data)
    return data, None


# ---------------------------------------------------------------- 연차 일정

TIME_RE = re.compile(r"(\d{1,2}):(\d{2})\s*~\s*(\d{1,2}):(\d{2})")


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    # 셀 서식이 깨지면 날짜가 엑셀 일련번호(예: 46325 = 2026-10-30)로 읽힌다
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 30000 < v < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=int(v))).date()
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d{4})[-./]\s*(\d{1,2})[-./]\s*(\d{1,2})\.?\s*", v)
        if m:
            return date(*map(int, m.groups()))
    return None


def hours_between(m):
    h1, m1, h2, m2 = map(int, m.groups())
    return max(0.0, (h2 * 60 + m2 - h1 * 60 - m1) / 60)


def classify(raw, hours_per_day):
    """'시간단위휴가(10:00~14:00), 대외활동(14:00~15:30)' 같은 구분값을 해석한다.

    반환: (category, per_day_deduction, trip, label)
      category: 연차 | 오전반차 | 오후반차 | 시간단위 | 공가 | 출장 | 기타
    """
    text = str(raw).strip()
    deduct, trip, cats = 0.0, False, []
    for part in [s.strip() for s in text.split(",") if s.strip()]:
        tm = TIME_RE.search(part)
        if part.startswith("연차"):
            cats.append("연차"); deduct += 1
        elif "오전반차" in part:
            cats.append("오전반차"); deduct += 0.5
        elif "오후반차" in part:
            cats.append("오후반차"); deduct += 0.5
        elif part.startswith("시간단위") or part.startswith("시간"):
            cats.append("시간단위")
            if tm:
                deduct += hours_between(tm) / hours_per_day
        elif part.startswith("공가") or part.startswith("병가") or part.startswith("경조"):
            cats.append("공가")
        elif "출장" in part or "대외활동" in part or "외근" in part or "교육" in part:
            cats.append("출장"); trip = True
        else:
            cats.append("기타")
    order = ["연차", "오전반차", "오후반차", "공가", "시간단위", "출장", "기타"]
    category = next((c for c in order if c in cats), "기타")
    return category, round(deduct, 4), trip, text


def business_days(start, end, holidays):
    d, out = start, []
    while d <= end:
        if d.weekday() < 5 and d.isoformat() not in holidays:
            out.append(d)
        d += timedelta(days=1)
    return out


def parse_leave(path, sheet, hours_per_day):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = list(r) + [None] * 5
        name, s, e, kind, note = r[0], to_date(r[1]), to_date(r[2]), r[3], r[4]
        if not name or not s:
            continue
        rows.append((i, str(name).strip(), s, e or s, kind, note))
    wb.close()

    holidays = {}
    for _, name, s, e, kind, _n in rows:
        if kind in (None, ""):
            d = s
            while d <= e:
                holidays[d.isoformat()] = name
                d += timedelta(days=1)

    entries = []
    for row, name, s, e, kind, note in rows:
        if kind in (None, ""):
            continue
        category, per_day, trip, label = classify(kind, hours_per_day)
        days = business_days(s, e, holidays)
        entries.append({
            "id": row,
            "name": name,
            "start": s.isoformat(),
            "end": e.isoformat(),
            "category": category,
            "label": label,
            "trip": trip,
            "note": "" if note is None else str(note),
            "workdays": len(days),
            "deduct": round(per_day * len(days), 4),
            "dates": [d.isoformat() for d in days],
        })
    return {"holidays": holidays, "entries": entries}


# ---------------------------------------------------------------- 주간업무계획

def norm(s):
    return unicodedata.normalize("NFKC", str(s)) if s is not None else ""


def clean_text(v):
    if v is None:
        return ""
    lines = []
    for ln in norm(v).replace("\r", "").split("\n"):
        ln = ln.rstrip()
        if ln.startswith("'") or ln.startswith("`-"):
            ln = ln[1:]
        lines.append(ln)
    return "\n".join(lines).strip()


SHEET_RE = re.compile(r"(\d{2})년\s*(\d{1,2})월\s*([\d·,]+)주차_?(.*)")


def parse_weekly(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    weeks = []
    for ws in wb.worksheets:
        title = norm(ws.title)
        m = SHEET_RE.match(title)
        header, items, blank = None, [], 0
        for r in ws.iter_rows(min_row=1, max_col=6, values_only=True):
            r = list(r) + [None] * 6
            if header is None:
                header = [clean_text(c) for c in r[:6]]
                continue
            if not any(c not in (None, "") for c in r[:6]):
                blank += 1
                if blank > 20:
                    break
                continue
            blank = 0
            area = clean_text(r[0])
            if not area:
                continue
            author_raw = clean_text(r[4])
            author, written = author_raw, ""
            if "/" in author_raw:
                author, written = [x.strip() for x in author_raw.split("/", 1)]
            items.append({
                "area": area,
                "progress": clean_text(r[1]),
                "plan": clean_text(r[2]),
                "key": clean_text(r[3]),
                "author": author,
                "written": written,
                "reviewed": clean_text(r[5]).upper() == "Y",
            })
        if header is None:
            continue
        weeks.append({
            "sheet": ws.title,
            "title": title,
            "year": int("20" + m.group(1)) if m else None,
            "month": int(m.group(2)) if m else None,
            "week": m.group(3) if m else "",
            "status": (m.group(4) if m else "").strip("_ ") or "",
            "progress_label": header[1].replace("\n", " ").strip() or "진행",
            "plan_label": header[2].replace("\n", " ").strip() or "계획",
            "items": items,
        })
    wb.close()
    return weeks


# ---------------------------------------------------------------- 개인별 주간업무보고

PERIOD_RE = re.compile(r"\((\d{4})\.(\d{1,2})\.(\d{1,2})\s*~\s*(?:(\d{1,2})\.)?(\d{1,2})\)")
STATUS_MAP = {"진행중": "진행 중", "완료됨": "완료", "완료": "완료", "시작전": "시작 전", "중단됨": "중단", "중단": "중단"}


def norm_status(v):
    s = re.sub(r"\s+", "", norm(v))
    return STATUS_MAP.get(s, norm(v).strip() or "미지정")


def parse_personal(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    weeks = []
    for ws in wb.worksheets:
        title = norm(ws.title).replace("－", "-")
        m = PERIOD_RE.search(title)
        start = end = None
        if m:
            y, mo, d, emo, ed = m.groups()
            start = date(int(y), int(mo), int(d))
            end = date(int(y), int(emo or mo), int(ed))
            if end < start:  # 연도 넘김(12월~1월)
                end = date(int(y) + 1, end.month, end.day)
        tasks, header, blank = [], None, 0
        for i, r in enumerate(ws.iter_rows(max_col=11, values_only=True), start=1):
            r = list(r) + [None] * 11
            if header is None:
                header = r
                continue
            owner, name, content = (clean_text(x) for x in r[1:4])
            if not (owner or name or content):
                blank += 1
                if blank > 30:
                    break
                continue
            blank = 0
            tasks.append({
                "row": i,
                "owner": owner or "미지정",
                "name": name,
                "content": content,
                "priority": norm(r[4]).count("★"),
                "status": norm_status(r[5]),
                "start": d.isoformat() if (d := to_date(r[6])) else "",
                "due": d.isoformat() if (d := to_date(r[7])) else "",
                "done": d.isoformat() if (d := to_date(r[8])) else "",
                "result": clean_text(r[9]),
                "memo": clean_text(r[10]),
            })
        weeks.append({
            "sheet": ws.title,
            "title": PERIOD_RE.sub("", title).strip(),
            "start": start.isoformat() if start else "",
            "end": end.isoformat() if end else "",
            "tasks": tasks,
        })
    wb.close()
    weeks.sort(key=lambda w: w["start"], reverse=True)  # 최신 주차가 앞으로
    return weeks


# ---------------------------------------------------------------- 사전협의 접수·검토 현황

REVIEW_FIELDS = {
    "no": "협의번호", "kind": "구분", "agency": "신청기관명", "type": "신청유형", "field": "신청분야",
    "biz1": "세부사업명", "biz2": "내역사업명", "biz3": "내내역사업명", "state": "진행상태",
    "assignee": "검토담당자", "received": "접수완료", "reviewed": "검토완료", "result": "결과서",
    "memo": "메모", "draft": "초안작성", "draft_review": "초안검토", "sent": "중기청송부",
}
REVIEW_DATES = {"received", "reviewed", "draft", "draft_review", "sent"}
MD_RE = re.compile(r"\s*(\d{1,2})\s*(?:월\s*|[-./]\s*)(\d{1,2})\s*일?")


def review_date(v, year):
    """'9월 8일', '05-06', 날짜 셀을 ISO로 바꾼다. 해석할 수 없으면 원문을 그대로 둔다."""
    if v in (None, "", "-"):
        return ""
    d = to_date(v)
    if d:
        return d.isoformat()
    text = norm(v).strip()
    m = MD_RE.match(text)
    if m:
        try:
            return date(year, int(m.group(1)), int(m.group(2))).isoformat()
        except ValueError:
            pass
    return text


def latest_file(pattern):
    """'..._*.xlsx'처럼 날짜가 붙는 파일은 이름순으로 가장 최근 파일을 쓴다."""
    if not any(c in pattern for c in "*?["):
        return pattern
    files = [f for f in glob.glob(pattern) if not Path(f).name.startswith("~$")]
    return max(files, key=lambda f: Path(f).name) if files else pattern


def parse_reviews(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    cols, rows = None, []
    for r in ws.iter_rows(values_only=True):
        if cols is None:
            names = [norm(c).strip() for c in r]
            if "협의번호" in names:
                cols = {k: names.index(h) for k, h in REVIEW_FIELDS.items() if h in names}
            continue
        r = list(r) + [None] * len(REVIEW_FIELDS)
        no = norm(r[cols["no"]]).strip()
        if not no:
            continue
        year = int(no[:4]) if no[:4].isdigit() else date.today().year
        row = {}
        for k, i in cols.items():
            v = r[i]
            row[k] = review_date(v, year) if k in REVIEW_DATES else ("" if v in (None, "-") else clean_text(v))
        row["no"] = no
        rows.append(row)
    wb.close()
    if cols is None:
        raise ValueError("'협의번호' 헤더를 찾지 못했습니다.")
    rows.sort(key=lambda x: x["no"], reverse=True)
    return {"file": Path(path).name, "rows": rows}


# ---------------------------------------------------------------- 집계

def build_payload():
    cfg = load_config()
    hpd = cfg.get("hours_per_day", 8)
    errors = []
    leave, err = cached(cfg["leave_file"], lambda p: parse_leave(p, cfg.get("leave_sheet", "연차입력"), hpd))
    if err:
        errors.append(err)
    weekly, err = cached(cfg["weekly_file"], parse_weekly)
    if err:
        errors.append(err)
    personal = None
    if cfg.get("personal_file"):
        personal, err = cached(cfg["personal_file"], parse_personal)
        if err:
            errors.append(err)
    reviews = None
    if cfg.get("review_file"):
        reviews, err = cached(latest_file(cfg["review_file"]), parse_reviews)
        if err:
            errors.append(err)
    leave = leave or {"holidays": {}, "entries": []}

    today = date.today().isoformat()
    # 팀원은 config.json 명단 기준. 명단에 없는 이름(전출자 등)의 휴가·업무는 화면에서 제외한다.
    members = {m["name"]: {k: v for k, v in m.items() if k not in ("email", "admin")} for m in cfg.get("members", [])}
    entries = [e for e in leave["entries"] if e["name"] in members]
    personal = [dict(w, tasks=[t for t in w["tasks"] if t["owner"] in members]) for w in (personal or [])]
    for m in members.values():
        used = planned = 0.0
        for e in entries:
            if e["name"] != m["name"] or not e["deduct"]:
                continue
            per = e["deduct"] / max(len(e["dates"]), 1)
            for d in e["dates"]:
                if d <= today:
                    used += per
                else:
                    planned += per
        m["used"] = round(used + (m.get("prior_used") or 0), 2)
        m["planned"] = round(planned, 2)
        total = m.get("total")
        m["remaining"] = round(total - m["used"] - m["planned"], 2) if isinstance(total, (int, float)) else None

    return {
        "team_name": cfg.get("team_name", "팀"),
        "today": today,
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "hours_per_day": hpd,
        "members": list(members.values()),
        "holidays": leave["holidays"],
        "entries": entries,
        "weeks": weekly or [],
        "personal": personal,
        "reviews": reviews or {"file": "", "rows": []},
        "sources": {"leave": cfg["leave_file"], "weekly": cfg["weekly_file"], "personal": cfg.get("personal_file", ""),
                    "reviews": cfg.get("review_file", "")},
        "errors": errors,
    }


# ---------------------------------------------------------------- HTTP

class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw):
        super().__init__(*a, directory=str(STATIC), **kw)

    def log_message(self, fmt, *args):
        pass

    def end_headers(self):
        # 화면 파일을 고친 뒤 브라우저가 옛 캐시를 쓰지 않도록 매번 재검증하게 한다
        if "Cache-Control" not in "".join(h.decode("latin-1") for h in getattr(self, "_headers_buffer", [])):
            self.send_header("Cache-Control", "no-cache")
        super().end_headers()

    def send_json(self, obj, status=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/api/data":
            with _lock:
                return self.send_json(build_payload())
        return super().do_GET()


class Server(ThreadingHTTPServer):
    # Windows에서 SO_REUSEADDR를 켜면 같은 포트에 서버가 여러 개 떠서 옛 버전이 응답할 수 있다
    allow_reuse_address = False


def main():
    cfg = load_config()
    port = cfg.get("port", 8765)
    url = f"http://127.0.0.1:{port}/"
    try:
        server = Server(("127.0.0.1", port), Handler)
    except OSError:
        print(f"이미 HR관리앱이 실행 중입니다: {url}\n새 버전으로 다시 띄우려면 기존 HR앱 창을 닫고 다시 실행하세요.")
        webbrowser.open(url)
        return
    print(f"{cfg.get('team_name', '')} HR관리앱 실행 중: {url}  (종료: Ctrl+C)")
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
